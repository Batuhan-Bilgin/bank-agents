import os
import sys
import json
import hashlib
from pathlib import Path

DB_PATH = Path(__file__).parent / "db"
WORKFLOWS_PATH = Path(__file__).parent / "workflows"

SUPPORTED = {".txt", ".md", ".pdf", ".docx", ".xlsx"}


def _read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_pdf(path: Path) -> str:
    import pdfplumber
    text = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                text.append(t)
    return "\n".join(text)


def _read_docx(path: Path) -> str:
    from docx import Document
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _read_xlsx(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = []
    for sheet in wb.worksheets:
        rows.append(f"[Sayfa: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows.append(" | ".join(cells))
    return "\n".join(rows)


def read_document(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in (".txt", ".md"):
        return _read_txt(path)
    elif ext == ".pdf":
        return _read_pdf(path)
    elif ext == ".docx":
        return _read_docx(path)
    elif ext == ".xlsx":
        return _read_xlsx(path)
    return ""


def chunk_text(text: str, chunk_size: int = 600, overlap: int = 100) -> list[str]:
    words = text.split()
    chunks = []
    i = 0
    while i < len(words):
        chunk = " ".join(words[i:i + chunk_size])
        if chunk.strip():
            chunks.append(chunk)
        i += chunk_size - overlap
    return chunks


def _get_collection():
    import chromadb
    from chromadb.utils import embedding_functions
    client = chromadb.PersistentClient(path=str(DB_PATH))
    ef = embedding_functions.SentenceTransformerEmbeddingFunction(
        model_name="paraphrase-multilingual-MiniLM-L12-v2"
    )
    return client.get_or_create_collection(name="bank_workflows", embedding_function=ef)


def ingest_file(path: Path, domain: str, collection) -> int:
    print(f"  Okunuyor: {path.name}")
    text = read_document(path)
    if not text.strip():
        print(f"  [!] Boş döküman: {path.name}")
        return 0

    chunks = chunk_text(text)
    file_hash = hashlib.md5(path.read_bytes()).hexdigest()[:8]

    ids = [f"{file_hash}_{i}" for i in range(len(chunks))]
    metas = [{"source": path.name, "domain": domain, "chunk": i} for i in range(len(chunks))]

    existing = collection.get(ids=ids, include=[])["ids"]
    new_ids = [i for i in ids if i not in existing]
    new_chunks = [chunks[ids.index(i)] for i in new_ids]
    new_metas = [metas[ids.index(i)] for i in new_ids]

    if new_ids:
        collection.add(documents=new_chunks, ids=new_ids, metadatas=new_metas)
        print(f"  [+] {len(new_ids)} yeni chunk eklendi ({path.name})")
    else:
        print(f"  [=] Zaten güncel ({path.name})")

    return len(new_ids)


def ingest_all():
    if not WORKFLOWS_PATH.exists():
        print("training/workflows/ klasörü bulunamadı.")
        return

    collection = _get_collection()
    total = 0

    for item in sorted(WORKFLOWS_PATH.iterdir()):
        if item.is_dir():
            domain = item.name
            print(f"\n[Domain: {domain}]")
            for f in sorted(item.iterdir()):
                if f.suffix.lower() in SUPPORTED:
                    total += ingest_file(f, domain, collection)
        elif item.suffix.lower() in SUPPORTED:
            domain = item.stem.split("_")[0]
            print(f"\n[Domain: {domain}]")
            total += ingest_file(item, domain, collection)

    print(f"\nToplam {total} yeni chunk eklendi.")
    print(f"Veritabanı: {DB_PATH}")


def list_domains():
    collection = _get_collection()
    all_meta = collection.get(include=["metadatas"])["metadatas"]
    domains = {}
    for m in all_meta:
        d = m.get("domain", "?")
        s = m.get("source", "?")
        if d not in domains:
            domains[d] = set()
        domains[d].add(s)
    print(f"\nToplam {collection.count()} chunk, {len(domains)} domain:\n")
    for domain, sources in sorted(domains.items()):
        print(f"  [{domain}]")
        for s in sorted(sources):
            print(f"    - {s}")


def delete_domain(domain: str):
    collection = _get_collection()
    results = collection.get(where={"domain": domain}, include=["metadatas"])
    ids = results["ids"]
    if ids:
        collection.delete(ids=ids)
        print(f"{len(ids)} chunk silindi (domain: {domain})")
    else:
        print(f"Domain bulunamadı: {domain}")


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else "ingest"

    if cmd == "ingest":
        ingest_all()
    elif cmd == "list":
        list_domains()
    elif cmd == "delete" and len(sys.argv) > 2:
        delete_domain(sys.argv[2])
    else:
        print("Kullanim:")
        print("  python training/ingest.py ingest          # Tum dokumanlari isle")
        print("  python training/ingest.py list            # Mevcut domainleri listele")
        print("  python training/ingest.py delete <domain> # Domain sil")
